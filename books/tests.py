import io
from decimal import Decimal

from django.contrib.auth.models import Permission, User
from django.contrib.messages import get_messages
from django.core.exceptions import PermissionDenied, ValidationError
from django.test import TestCase
from openpyxl import load_workbook

from core.testing import make_account, make_project, make_request
from costing.forms import SharedCostForm
from costing.models import CostCategory, ProjectCost
from customers.models import Customer
from sales.models import CustomerPayment, FlatSale, HonorariumPayment, Reference, RefundPayment

from .forms import TransferForm
from .ledger import (
    account_summaries,
    assign_account,
    collect_entries,
    confirm_entries,
    lowest_balance,
    pending_count,
    unassigned_count,
    with_running_balance,
)
from .models import Account, Transfer
from .views import (
    AssignAccountView,
    BookView,
    ConfirmEntriesView,
    ConfirmPageView,
    TransferDeleteView,
    pick_account,
)

D = Decimal


class BookTestCase(TestCase):
    def setUp(self):
        self.project, self.flats = make_project()
        self.bank1 = make_account('Bank 1', opening_balance=10_000)
        self.bank2 = make_account('Bank 2')
        self.cash = make_account('Petty cash', account_type=Account.AccountType.CASH)
        self.customer = Customer.objects.create(name='Rahim', phone='01711111111')
        self.ref = Reference.objects.create(name='Ref')
        self.sale = FlatSale.objects.create(
            flat=self.flats[0], customer=self.customer, sale_date='2026-01-01', sale_value=1500, ref=self.ref,
        )
        self.category = CostCategory.objects.create(name='Civil')
        self.boss = User.objects.create_superuser('boss', 'b@example.com', 'pw')

    def pay(self, amount, account=None, date='2026-01-05'):
        return CustomerPayment.objects.create(sale=self.sale, payment_date=date, amount=amount, account=account or self.bank1)

    def spend(self, amount, account=None, date='2026-01-06'):
        return ProjectCost.objects.create(
            project=self.project, cost_category=self.category, date=date, amount=amount, account=account or self.bank1,
        )

    def confirm(self, *items, user=None):
        kinds = {CustomerPayment: 'cp', ProjectCost: 'pc', Transfer: 'tr', HonorariumPayment: 'hp', RefundPayment: 'rf'}
        return confirm_entries([f'{kinds[type(i)]}:{i.pk}' for i in items], user or self.boss)

    def confirmed(self, item):
        return type(item).objects.get(pk=item.pk).is_confirmed


class LedgerCollectionTests(BookTestCase):
    def test_every_source_shows_up_with_the_right_direction(self):
        self.pay(1000)
        self.spend(200)
        HonorariumPayment.objects.create(sale=self.sale, date='2026-01-07', amount=100, account=self.bank1)
        self.sale.cancel(50)
        self.sale.save()
        RefundPayment.objects.create(sale=self.sale, date='2026-01-08', amount=300, account=self.bank1)
        Transfer.objects.create(date='2026-01-09', from_account=self.bank1, to_account=self.cash, amount=50)

        by_source = {}
        for entry in collect_entries([self.bank1]):
            by_source.setdefault(entry.source, []).append(entry)
        self.assertEqual(by_source['Payment received'][0].money_in, 1000)
        self.assertEqual(by_source['Cost payment'][0].money_out, 200)
        self.assertEqual(by_source['Honorarium paid'][0].money_out, 100)
        self.assertEqual(by_source['Refund to customer'][0].money_out, 300)
        self.assertEqual(by_source['Transfer out'][0].money_out, 50)
        self.assertNotIn('Transfer in', by_source)

    def test_a_transfer_is_out_of_one_account_and_into_the_other(self):
        Transfer.objects.create(date='2026-01-09', from_account=self.bank1, to_account=self.cash, amount=500)
        out_side = collect_entries([self.bank1])
        in_side = collect_entries([self.cash])
        self.assertEqual((out_side[0].money_out, out_side[0].money_in), (500, 0))
        self.assertEqual((in_side[0].money_in, in_side[0].money_out), (500, 0))

    def test_transfer_between_the_same_account_is_rejected(self):
        form = TransferForm(data={'date': '2026-01-01', 'from_account': self.bank1.pk, 'to_account': self.bank1.pk, 'amount': '5'})
        self.assertFalse(form.is_valid())
        with self.assertRaises(ValidationError):
            Transfer(from_account=self.bank1, to_account=self.bank1, amount=1, date='2026-01-01').clean()

    def test_the_confirmed_filter(self):
        first = self.pay(100)
        self.pay(200)
        self.confirm(first)
        self.assertEqual([e.money_in for e in collect_entries([self.bank1], confirmed=True)], [100])
        self.assertEqual([e.money_in for e in collect_entries([self.bank1], confirmed=False)], [200])
        self.assertEqual(len(collect_entries([self.bank1])), 2)

    def test_the_book_holds_no_data_of_its_own(self):
        self.pay(1000)
        before = len(collect_entries())
        self.pay(500)
        self.assertEqual(len(collect_entries()), before + 1)
        CustomerPayment.objects.first().delete()
        self.assertEqual(len(collect_entries()), before)

    def test_records_without_an_account_are_reported_not_counted(self):
        CustomerPayment.objects.create(sale=self.sale, payment_date='2026-01-05', amount=77)
        ProjectCost.objects.create(project=self.project, cost_category=self.category, date='2026-01-05', amount=11)
        self.assertEqual(unassigned_count(), 2)
        self.assertEqual(collect_entries(), [])
        self.assertEqual(sum(r['balance'] for r in account_summaries(Account.objects.all())), 10_000)

    def test_running_balance_starts_from_the_opening_balance(self):
        self.pay(1000, date='2026-01-02')
        self.spend(400, date='2026-01-03')
        entries = collect_entries([self.bank1])
        closing = with_running_balance(entries, self.bank1.opening_balance)
        self.assertEqual([e.balance for e in entries], [11_000, 10_600])
        self.assertEqual(closing, 10_600)


class ConfirmationTests(BookTestCase):
    def test_new_entries_are_not_in_the_book_until_confirmed(self):
        payment = self.pay(1000)
        self.assertEqual(collect_entries([self.bank1], confirmed=True), [])
        summary = account_summaries([self.bank1])[0]
        self.assertEqual((summary['balance'], summary['pending_count']), (10_000, 1))
        done, problems = self.confirm(payment)
        self.assertEqual((done, problems), ([f'cp:{payment.pk}'], []))
        summary = account_summaries([self.bank1])[0]
        self.assertEqual((summary['balance'], summary['pending_count']), (11_000, 0))

    def test_confirming_records_who_and_when(self):
        payment = self.pay(1000)
        self.confirm(payment)
        payment.refresh_from_db()
        self.assertEqual(payment.confirmed_by, self.boss)
        self.assertIsNotNone(payment.confirmed_at)

    def test_a_confirmation_cannot_be_undone(self):
        payment = self.pay(1000)
        self.confirm(payment)
        self.assertFalse(hasattr(payment, 'unconfirm'))
        # an old-style "unconfirm" request changes nothing
        request = make_request(self.boss, '/books/confirm/', 'post', {'entries': [f'cp:{payment.pk}'], 'action': 'unconfirm'})
        ConfirmEntriesView.as_view()(request)
        self.assertTrue(self.confirmed(payment))

    def test_confirming_again_is_harmless(self):
        payment = self.pay(1000)
        self.confirm(payment)
        done, problems = self.confirm(payment)
        self.assertEqual(done, [])
        self.assertEqual(len(problems), 1)
        self.assertEqual(account_summaries([self.bank1])[0]['balance'], 11_000)

    def test_confirming_does_not_disturb_the_sale_totals(self):
        payment = self.pay(1000)
        before = FlatSale.objects.get(pk=self.sale.pk).received_amount
        self.confirm(payment)
        self.assertEqual(FlatSale.objects.get(pk=self.sale.pk).received_amount, before)

    def test_every_kind_of_entry_can_be_confirmed(self):
        self.pay(5000)
        payment = CustomerPayment.objects.get()
        cost = self.spend(300)
        transfer = Transfer.objects.create(date='2026-01-09', from_account=self.bank1, to_account=self.cash, amount=5)
        honorarium = HonorariumPayment.objects.create(sale=self.sale, date='2026-01-07', amount=100, account=self.bank1)
        done, problems = self.confirm(payment, cost, transfer, honorarium)
        self.assertEqual((len(done), problems), (4, []))

    def test_entries_without_an_account_cannot_be_confirmed(self):
        payment = CustomerPayment.objects.create(sale=self.sale, payment_date='2026-01-05', amount=77)
        done, problems = self.confirm(payment)
        self.assertEqual((done, len(problems)), ([], 1))
        self.assertFalse(self.confirmed(payment))


class NeverNegativeTests(BookTestCase):
    """Bank 1 starts with 10,000."""

    def test_a_payment_out_that_overdraws_the_account_is_refused(self):
        cost = self.spend(15_000)
        done, problems = self.confirm(cost)
        self.assertEqual(done, [])
        self.assertIn('Bank 1', problems[0])
        self.assertIn('below zero', problems[0])
        self.assertFalse(self.confirmed(cost))
        self.assertEqual(account_summaries([self.bank1])[0]['balance'], 10_000)

    def test_it_is_fine_once_enough_money_has_come_in(self):
        payment = self.pay(6_000, date='2026-01-02')
        cost = self.spend(15_000, date='2026-01-06')
        done, problems = self.confirm(cost, payment)   # order of ticking does not matter: oldest goes first
        self.assertEqual((len(done), problems), (2, []))
        self.assertEqual(account_summaries([self.bank1])[0]['balance'], 1_000)

    def test_the_balance_may_reach_exactly_zero(self):
        cost = self.spend(10_000)
        self.assertEqual(self.confirm(cost)[1], [])
        self.assertEqual(account_summaries([self.bank1])[0]['balance'], 0)

    def test_a_backdated_payment_out_cannot_dip_below_zero_in_between(self):
        late_payment = self.pay(6_000, date='2026-01-10')
        self.confirm(late_payment)
        early_cost = self.spend(15_000, date='2026-01-06')   # would be negative on 6 Jan, before the money arrived
        done, problems = self.confirm(early_cost)
        self.assertEqual(done, [])
        self.assertFalse(self.confirmed(early_cost))

    def test_one_refused_entry_does_not_block_the_others(self):
        small = self.spend(100, date='2026-01-03')
        huge = self.spend(50_000, date='2026-01-04')
        done, problems = self.confirm(small, huge)
        self.assertEqual(done, [f'pc:{small.pk}'])
        self.assertEqual(len(problems), 1)

    def test_a_transfer_needs_the_money_in_the_account_it_leaves(self):
        too_much = Transfer.objects.create(date='2026-01-09', from_account=self.bank2, to_account=self.cash, amount=1)
        done, problems = self.confirm(too_much)
        self.assertEqual(done, [])
        self.assertIn('Bank 2', problems[0])
        self.assertEqual(lowest_balance(self.cash), 0)

    def test_a_transfer_moves_money_between_confirmed_balances(self):
        transfer = Transfer.objects.create(date='2026-01-09', from_account=self.bank1, to_account=self.cash, amount=4_000)
        self.assertEqual(self.confirm(transfer)[1], [])
        balances = {r['account'].name: r['balance'] for r in account_summaries(Account.objects.all())}
        self.assertEqual((balances['Bank 1'], balances['Petty cash']), (6_000, 4_000))

    def test_refunds_and_honorarium_are_money_out_too(self):
        honorarium = HonorariumPayment.objects.create(sale=self.sale, date='2026-01-07', amount=10, account=self.bank2)
        self.assertEqual(self.confirm(honorarium)[0], [])
        self.assertFalse(self.confirmed(honorarium))


class ConfirmViewsTests(BookTestCase):
    def post(self, user, keys, view=ConfirmEntriesView, **extra):
        request = make_request(user, '/books/confirm/', 'post', {'entries': keys, **extra})
        response = view.as_view()(request)
        return response, [str(m) for m in get_messages(request)]

    def test_confirming_needs_the_permission(self):
        payment = self.pay(1000)
        clerk = User.objects.create_user('clerk')
        with self.assertRaises(PermissionDenied):
            self.post(clerk, [f'cp:{payment.pk}'])
        clerk.user_permissions.add(Permission.objects.get(codename='confirm_entries'))
        self.post(User.objects.get(pk=clerk.pk), [f'cp:{payment.pk}'])
        self.assertTrue(self.confirmed(payment))

    def test_the_view_reports_what_was_refused(self):
        cost = self.spend(15_000)
        response, messages_ = self.post(self.boss, [f'pc:{cost.pk}'])
        self.assertEqual(response.status_code, 302)
        self.assertTrue(any('below zero' in m for m in messages_))

    def test_the_view_says_so_when_it_is_final(self):
        payment = self.pay(1000)
        _, messages_ = self.post(self.boss, [f'cp:{payment.pk}'])
        self.assertTrue(any('cannot be undone' in m for m in messages_))

    def test_bad_keys_are_ignored_not_crashed_on(self):
        response, _ = self.post(self.boss, ['zz:1', 'cp:abc', 'cp:99999', 'nonsense'])
        self.assertEqual(response.status_code, 302)

    def test_external_next_urls_are_ignored(self):
        response, _ = self.post(self.boss, [], next='https://evil.example.com/')
        self.assertEqual(response['Location'], '/books/confirm-entries/')

    def test_the_confirm_page_shows_one_account_and_its_waiting_entries(self):
        self.pay(100, self.bank1)
        self.pay(200, self.bank2)
        view = ConfirmPageView()
        view.setup(make_request(self.boss, '/books/confirm-entries/', data={'account': self.bank2.pk}))
        context = view.get_context_data()
        self.assertEqual(context['account'], self.bank2)
        self.assertEqual([e.money_in for e in context['entries']], [200])
        self.assertEqual(context['pending'][self.bank1.pk], 1)

    def test_the_confirm_page_opens_on_an_account_with_something_waiting(self):
        self.pay(200, self.bank2)
        view = ConfirmPageView()
        view.setup(make_request(self.boss, '/books/confirm-entries/'))
        self.assertEqual(view.get_context_data()['account'], self.bank2)

    def test_the_confirm_page_needs_the_permission(self):
        clerk = User.objects.create_user('clerk')
        with self.assertRaises(PermissionDenied):
            ConfirmPageView.as_view()(make_request(clerk, '/books/confirm-entries/'))

    def test_a_confirmed_transfer_cannot_be_deleted(self):
        transfer = Transfer.objects.create(date='2026-01-09', from_account=self.bank1, to_account=self.cash, amount=5)
        self.confirm(transfer)
        response = TransferDeleteView.as_view()(make_request(self.boss, '/', 'post'), pk=transfer.pk)
        self.assertEqual(response.status_code, 302)
        self.assertTrue(Transfer.objects.filter(pk=transfer.pk).exists())

    def test_an_unconfirmed_transfer_can_be_deleted(self):
        transfer = Transfer.objects.create(date='2026-01-09', from_account=self.bank1, to_account=self.cash, amount=5)
        TransferDeleteView.as_view()(make_request(self.boss, '/', 'post'), pk=transfer.pk)
        self.assertFalse(Transfer.objects.filter(pk=transfer.pk).exists())

    def test_menu_badge_counts_what_waits(self):
        self.pay(100)
        self.spend(5)
        self.assertEqual(pending_count(), 2)
        self.confirm(CustomerPayment.objects.get())
        self.assertEqual(pending_count(), 1)


class AssignAccountTests(BookTestCase):
    def test_older_records_get_an_account_and_then_wait_for_confirmation(self):
        payment = CustomerPayment.objects.create(sale=self.sale, payment_date='2026-01-05', amount=77)
        cost = ProjectCost.objects.create(project=self.project, cost_category=self.category, date='2026-01-05', amount=11)
        self.assertEqual(assign_account([f'cp:{payment.pk}', f'pc:{cost.pk}'], self.cash), 2)
        payment.refresh_from_db()
        cost.refresh_from_db()
        self.assertEqual((payment.account, cost.account, payment.payment_method), (self.cash, self.cash, 'CASH'))
        self.assertEqual(unassigned_count(), 0)
        self.assertEqual(len(collect_entries([self.cash], confirmed=False)), 2)

    def test_records_that_already_have_an_account_are_left_alone(self):
        payment = self.pay(100)
        self.assertEqual(assign_account([f'cp:{payment.pk}'], self.cash), 0)
        payment.refresh_from_db()
        self.assertEqual(payment.account, self.bank1)

    def test_other_kinds_of_entry_cannot_be_reassigned(self):
        transfer = Transfer.objects.create(date='2026-01-09', from_account=self.bank1, to_account=self.cash, amount=5)
        self.assertEqual(assign_account([f'tr:{transfer.pk}', 'zz:1', 'cp:x'], self.bank2), 0)

    def test_the_view(self):
        payment = CustomerPayment.objects.create(sale=self.sale, payment_date='2026-01-05', amount=77)
        request = make_request(self.boss, '/', 'post', {'entries': [f'cp:{payment.pk}'], 'account': self.bank2.pk})
        self.assertEqual(AssignAccountView.as_view()(request).status_code, 302)
        payment.refresh_from_db()
        self.assertEqual(payment.account, self.bank2)

    def test_the_view_needs_the_permission(self):
        with self.assertRaises(PermissionDenied):
            AssignAccountView.as_view()(make_request(User.objects.create_user('clerk'), '/', 'post', {}))


class ConfirmedCostsAreLockedTests(BookTestCase):
    def test_a_confirmed_cost_cannot_be_edited(self):
        cost = self.spend(300)
        self.confirm(cost)
        cost.refresh_from_db()
        form = SharedCostForm(data={
            'project': self.project.pk, 'cost_category': self.category.pk, 'date': '2026-01-06', 'amount': '9999',
            'account': self.bank1.pk,
        }, instance=cost)
        self.assertFalse(form.is_valid())
        self.assertIn('confirmed', str(form.errors))

    def test_the_edit_and_delete_screens_refuse(self):
        from costing.views import ProjectCostDeleteView, ProjectCostUpdateView
        cost = self.spend(300)
        self.confirm(cost)
        for view in (ProjectCostUpdateView, ProjectCostDeleteView):
            response = view.as_view()(make_request(self.boss, '/', 'post', {}), pk=cost.pk)
            self.assertEqual(response.status_code, 302)
            self.assertEqual(response['Location'], cost.get_absolute_url())
        self.assertTrue(ProjectCost.objects.filter(pk=cost.pk, amount=300).exists())

    def test_an_unconfirmed_cost_is_still_editable_and_deletable(self):
        from costing.views import ProjectCostDeleteView
        cost = self.spend(300)
        ProjectCostDeleteView.as_view()(make_request(self.boss, '/', 'post', {}), pk=cost.pk)
        self.assertFalse(ProjectCost.objects.filter(pk=cost.pk).exists())


class BookViewTests(BookTestCase):
    def view(self, user=None, pk=None, **params):
        view = BookView()
        view.setup(make_request(user or self.boss, '/books/', data=params))
        view.kwargs = {'pk': pk} if pk else {}
        return view

    def context(self, **params):
        return self.view(**params).book()

    def test_the_book_shows_one_account_at_a_time(self):
        self.pay(100, self.bank1, date='2026-01-05')
        self.pay(200, self.bank2, date='2026-01-05')
        self.confirm(*CustomerPayment.objects.all())
        context = self.context(account=self.bank2.pk, date_from='2026-01-01', date_to='2026-01-31')
        self.assertEqual(context['account'], self.bank2)
        self.assertEqual({e.account for e in context['entries']}, {self.bank2})
        self.assertEqual([e.money_in for e in context['entries']], [200])

    def test_with_no_choice_it_opens_the_first_active_account(self):
        self.assertEqual(self.context()['account'], self.bank1)
        self.bank1.is_active = False
        self.bank1.save()
        self.assertEqual(self.context()['account'], self.bank2)

    def test_a_wrong_account_id_falls_back_instead_of_failing(self):
        self.assertEqual(self.context(account='abc')['account'], self.bank1)
        self.assertEqual(self.context(account='9999')['account'], self.bank1)
        self.assertEqual(pick_account([], 1), None)

    def test_only_confirmed_entries_are_in_the_book(self):
        confirmed = self.pay(100, date='2026-01-05')
        self.pay(200, date='2026-01-05')
        self.confirm(confirmed)
        context = self.context(account=self.bank1.pk, date_from='2026-01-01', date_to='2026-01-31')
        self.assertEqual([e.money_in for e in context['entries']], [100])
        self.assertEqual(context['closing_balance'], 10_100)

    def test_period_totals_and_opening_balance(self):
        january, february = self.pay(1000, date='2026-01-05'), self.spend(250, date='2026-02-05')
        self.confirm(january, february)
        context = self.context(account=self.bank1.pk, date_from='2026-02-01', date_to='2026-02-28')
        self.assertEqual(context['opening_balance'], 11_000)
        self.assertEqual((context['total_in'], context['total_out']), (0, 250))
        self.assertEqual(context['closing_balance'], 10_750)

    def test_the_closing_balance_is_never_negative(self):
        # build a messy history through the one door that exists: confirmation
        entries = [self.spend(a, date=f'2026-01-{d:02d}') for d, a in ((2, 4_000), (3, 4_000), (4, 4_000), (5, 4_000))]
        self.pay(7_000, date='2026-01-04')
        self.confirm(*entries, *CustomerPayment.objects.all())
        context = self.context(account=self.bank1.pk, date_from='2026-01-01', date_to='2026-12-31')
        self.assertGreaterEqual(context['closing_balance'], 0)
        self.assertTrue(all(e.balance >= 0 for e in context['entries']))

    def test_download_as_excel(self):
        payment = self.pay(1000, date='2026-01-05')
        self.confirm(payment)
        view = self.view(account=self.bank1.pk, date_from='2026-01-01', date_to='2026-01-31', format='xlsx')
        response = view.get(view.request)
        self.assertIn('spreadsheetml', response['Content-Type'])
        self.assertIn('attachment', response['Content-Disposition'])
        sheet = load_workbook(io.BytesIO(response.content)).active
        values = [[c.value for c in row] for row in sheet.iter_rows()]
        flat = [v for row in values for v in row]
        self.assertIn('Opening balance', flat)
        self.assertIn(1000.0, flat)       # real numbers, not text
        self.assertIn(11000.0, flat)      # closing balance

    def test_excel_only_has_confirmed_entries(self):
        self.pay(777, date='2026-01-05')
        view = self.view(account=self.bank1.pk, date_from='2026-01-01', date_to='2026-01-31', format='xlsx')
        sheet = load_workbook(io.BytesIO(view.get(view.request).content)).active
        self.assertNotIn(777.0, [c.value for row in sheet.iter_rows() for c in row])


class AccountTests(TestCase):
    def test_cash_and_bank_accounts(self):
        cash = make_account('Cash box', account_type=Account.AccountType.CASH)
        bank = make_account('Bank X', bank_name='X', account_number='123')
        self.assertTrue(cash.is_cash)
        self.assertEqual((cash.detail, bank.detail), ('Cash', 'X · 123'))

    def test_a_fresh_install_already_has_a_petty_cash_account(self):
        self.assertTrue(Account.objects.filter(name='Petty cash', account_type='CASH').exists())
