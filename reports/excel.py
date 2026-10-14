"""Excel (.xlsx) export of the same tables that are shown on screen and written to the PDFs."""
import datetime
import re
from decimal import Decimal

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.formatting import DATE, INT, MONEY, PERCENT

XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

HEADER_FILL = PatternFill('solid', fgColor='0B3D3A')
TOTAL_FILL = PatternFill('solid', fgColor='E2ECEA')
NUMBER_FORMATS = {MONEY: '#,##0.00', INT: '#,##0', PERCENT: '0.0%', DATE: 'dd mmm yyyy'}


def _sheet_name(title):
    return re.sub(r'[\[\]:*?/\\]', '-', title)[:31] or 'Report'


def _cell_value(value, kind):
    """A value ready for a spreadsheet cell: real numbers and dates, never formatted text."""
    if value is None or value == '':
        return None
    if kind == PERCENT:
        return float(Decimal(value) / 100)
    if kind in (MONEY, INT):
        return float(value) if isinstance(value, Decimal) else value
    if kind == DATE and isinstance(value, datetime.datetime):
        return value.date()
    return value


def build_xlsx(target, title, filter_lines, tables, user):
    """Write the report tables to `target` (file-like or HttpResponse) as one worksheet, stacked top to bottom."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _sheet_name(title)

    sheet.append([title])
    sheet['A1'].font = Font(bold=True, size=14, color='0B3D3A')
    for line in filter_lines:
        sheet.append([line])
    sheet.append([f'Generated {timezone.localtime():%d %b %Y, %I:%M %p} by {user.get_full_name() or user.username}'])
    widths = {}

    def remember(index, value):
        widths[index] = max(widths.get(index, 8), min(len(str(value)) + 2, 60))

    for table in tables:
        sheet.append([])
        sheet.append([table.title])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True, size=12)

        header_row = sheet.max_row + 1
        for index, column in enumerate(table.columns, start=1):
            cell = sheet.cell(row=header_row, column=index, value=column.label)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal='right' if column.numeric else 'left', wrap_text=True)
            remember(index, column.label)

        def write(row, bold=False):
            row_number = sheet.max_row + 1
            for index, (value, column) in enumerate(zip(row, table.columns), start=1):
                cell = sheet.cell(row=row_number, column=index, value=_cell_value(value, column.kind))
                if column.kind in NUMBER_FORMATS:
                    cell.number_format = NUMBER_FORMATS[column.kind]
                if bold:
                    cell.font = Font(bold=True)
                    cell.fill = TOTAL_FILL
                remember(index, value if value is not None else '')

        for row in table.rows:
            write(row)
        if table.totals:
            write(table.totals, bold=True)
        if table.note:
            sheet.append([table.note])
            sheet.cell(row=sheet.max_row, column=1).font = Font(italic=True, color='666666')

    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width
    workbook.save(target)
